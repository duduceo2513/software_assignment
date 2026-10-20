"""Builds one editable .docx from OCR page results (real text and real Word tables, no page images)."""
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from .ocr_engine import Block, PageResult

_XML_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")

HEADING_LABELS = {"doc_title": 0, "paragraph_title": 2}
CAPTION_LABELS = {"figure_title", "table_title", "chart_title", "caption"}
SKIPPED_LABELS = {"number", "header_image", "footer_image", "image", "seal"}
LINE_PRESERVING_LABELS = {"list", "reference", "algorithm"}


def clean(text: str) -> str:
    return _XML_ILLEGAL.sub("", text)


def _span(value: str | None) -> int:
    try:
        return max(1, min(int(value or 1), 50))
    except ValueError:
        return 1


class _TableParser(HTMLParser):
    """Minimal HTML table reader producing rows of (text, rowspan, colspan)."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[tuple[str, int, int]]] = []
        self._cell: dict | None = None
        self._depth = 0

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        if tag == "table":
            self._depth += 1
        elif self._depth != 1:
            return
        elif tag == "tr":
            self.rows.append([])
        elif tag in {"td", "th"}:
            if not self.rows:
                self.rows.append([])
            self._cell = {
                "text": [],
                "rowspan": _span(attributes.get("rowspan")),
                "colspan": _span(attributes.get("colspan")),
            }
        elif tag == "br" and self._cell is not None:
            self._cell["text"].append("\n")

    def handle_endtag(self, tag):
        if tag == "table":
            self._depth -= 1
        elif tag in {"td", "th"} and self._cell is not None and self._depth == 1:
            text = "".join(self._cell["text"]).strip()
            self.rows[-1].append((text, self._cell["rowspan"], self._cell["colspan"]))
            self._cell = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell["text"].append(data)


def parse_html_table(html: str) -> list[list[tuple[str, int, int]]]:
    parser = _TableParser()
    parser.feed(html)
    return [row for row in parser.rows if row]


def add_table(document, html: str) -> bool:
    """Add the HTML table to the document as a real Word table. Returns False if it has no cells."""
    rows = parse_html_table(html)
    if not rows:
        return False

    # Resolve rowspan/colspan into an absolute grid.
    placements: list[tuple[int, int, int, int, str]] = []
    occupied: set[tuple[int, int]] = set()
    n_cols = 0
    for r, row in enumerate(rows):
        c = 0
        for text, rowspan, colspan in row:
            while (r, c) in occupied:
                c += 1
            for dr in range(rowspan):
                for dc in range(colspan):
                    occupied.add((r + dr, c + dc))
            placements.append((r, c, rowspan, colspan, text))
            c += colspan
            n_cols = max(n_cols, c)
    n_rows = max(r for r, _ in occupied) + 1

    table = document.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    for r, c, rowspan, colspan, text in placements:
        cell = table.cell(r, c)
        if rowspan > 1 or colspan > 1:
            cell = cell.merge(table.cell(min(r + rowspan - 1, n_rows - 1), min(c + colspan - 1, n_cols - 1)))
        cell.text = clean(text)
    return True


def _add_text_block(document, block: Block) -> None:
    text = clean(block.text).strip()
    if not text:
        return
    if block.label in HEADING_LABELS:
        document.add_heading(" ".join(text.split()), level=HEADING_LABELS[block.label])
    elif block.label in CAPTION_LABELS:
        run = document.add_paragraph().add_run(" ".join(text.split()))
        run.italic = True
    elif block.label in LINE_PRESERVING_LABELS:
        for line in filter(None, (ln.strip() for ln in text.splitlines())):
            document.add_paragraph(line)
    else:
        document.add_paragraph(" ".join(text.split()))


def block_plain_text(block: Block) -> str:
    if block.table_html:
        return "\n".join(text for row in parse_html_table(block.table_html) for text, _, _ in row if text)
    return block.text


def export_docx(pages: list[PageResult], output_path: Path, title: str = "") -> str:
    """Write all pages, in order, to a single .docx. Returns the recognised plain text."""
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    if title:
        document.core_properties.title = clean(title)[:200]

    for position, page in enumerate(pages):
        if position:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        for block in page.blocks:
            if block.label in SKIPPED_LABELS and not block.text:
                continue
            if block.table_html:
                try:
                    if add_table(document, block.table_html):
                        continue
                except Exception:  # noqa: BLE001 - fall back to text so nothing is lost
                    pass
                _add_text_block(document, Block(block.label, block_plain_text(block)))
                continue
            _add_text_block(document, block)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)

    parts = [
        "\n".join(filter(None, (block_plain_text(b).strip() for b in page.blocks if b.label != "number")))
        for page in pages
    ]
    return "\n\n".join(part for part in parts if part)
