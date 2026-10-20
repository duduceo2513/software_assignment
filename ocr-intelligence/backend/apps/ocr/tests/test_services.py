import zipfile

from docx import Document

from apps.ocr.services.document_processor import iter_page_images
from apps.ocr.services.file_validator import sanitize_filename
from apps.ocr.services.ocr_engine import Block, PageResult, parse_structure_result, recognition_model_for
from apps.ocr.services.word_exporter import export_docx, parse_html_table
from apps.ocr.tests.conftest import make_pdf, make_png, make_tiff

KAZAKH = "Ә ә, Ғ ғ, Қ қ, Ң ң, Ө ө, Ұ ұ, Ү ү, Һ һ, І і"


def test_export_docx_is_editable_multipage_with_real_tables(tmp_path):
    html = (
        "<html><body><table>"
        "<tr><th colspan='2'>Итого</th></tr>"
        "<tr><td rowspan='2'>Қала</td><td>Алматы</td></tr>"
        "<tr><td>Астана</td></tr>"
        "</table></body></html>"
    )
    pages = [
        PageResult(1, [Block("doc_title", "Заголовок"), Block("text", f"Текст\nна двух строках {KAZAKH}")]),
        PageResult(2, [Block("paragraph_title", "Раздел 2"), Block("table", "", table_html=html)]),
    ]
    out = tmp_path / "out.docx"
    text = export_docx(pages, out, title="t")

    document = Document(out)
    assert not document.inline_shapes  # no page images embedded
    paragraphs = [p.text for p in document.paragraphs]
    assert "Заголовок" in paragraphs
    assert f"Текст на двух строках {KAZAKH}" in paragraphs
    assert paragraphs.index("Заголовок") < paragraphs.index("Раздел 2")
    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.rows) == 3
    assert table.cell(1, 1).text == "Алматы" and table.cell(2, 1).text == "Астана"
    assert table.cell(1, 0).text == "Қала"
    assert KAZAKH in text and "Астана" in text
    assert zipfile.is_zipfile(out)


def test_export_strips_xml_illegal_characters(tmp_path):
    out = tmp_path / "x.docx"
    export_docx([PageResult(1, [Block("text", "bad\x00char\x0b")])], out)
    assert Document(out).paragraphs[0].text == "badchar"


def test_broken_table_html_falls_back_to_text(tmp_path):
    out = tmp_path / "x.docx"
    export_docx([PageResult(1, [Block("table", "", table_html="<table></table>")])], out)
    assert Document(out).tables == []


def test_parse_html_table_spans():
    rows = parse_html_table("<table><tr><td colspan=2>a</td></tr><tr><td>b</td><td>c</td></tr></table>")
    assert rows == [[("a", 1, 2)], [("b", 1, 1), ("c", 1, 1)]]


def test_parse_structure_result_variants():
    result = {
        "res": {
            "parsing_res_list": [
                {"block_label": "doc_title", "block_content": "Title", "block_bbox": [0, 0, 1, 1]},
                {"block_label": "table", "block_content": "<table><tr><td>x</td></tr></table>"},
                {"block_label": "table", "block_content": ""},
                {"block_label": "text", "block_content": "  "},
            ],
            "table_res_list": [{"pred_html": "<table><tr><td>x</td></tr></table>"}, {"pred_html": "<table><tr><td>y</td></tr></table>"}],
        }
    }
    blocks = parse_structure_result(result)
    assert [b.label for b in blocks] == ["doc_title", "table", "table"]
    assert blocks[2].table_html == "<table><tr><td>y</td></tr></table>"

    fallback = parse_structure_result({"overall_ocr_res": {"rec_texts": ["a", " ", "b"]}})
    assert [b.text for b in fallback] == ["a", "b"]


def test_model_selection(settings):
    assert recognition_model_for("en") == settings.OCR_ENGLISH_REC_MODEL
    for language in ("auto", "kk", "ru"):
        assert recognition_model_for(language) == "cyrillic_PP-OCRv5_mobile_rec"


def test_sanitize_filename():
    assert sanitize_filename("..\\..\\win\\отчёт.pdf") == "отчёт.pdf"
    assert sanitize_filename("a<b>:c|?.png") == "a_b__c__.png"
    assert sanitize_filename("") == "document"


def test_iter_pages_pdf_and_tiff_and_png(tmp_path):
    for name, data, mime, expected in [
        ("a.pdf", make_pdf(3), "application/pdf", 3),
        ("b.tif", make_tiff(2), "image/tiff", 2),
        ("c.png", make_png(), "image/png", 1),
    ]:
        source = tmp_path / name
        source.write_bytes(data)
        pages = list(iter_page_images(source, mime, tmp_path / f"out-{name}"))
        assert [n for n, _ in pages] == list(range(1, expected + 1))
        assert all(path.exists() for _, path in pages)
