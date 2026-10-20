import io
import zipfile

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

from apps.ocr.models import DocumentStatus, JobStatus, OCRJob
from apps.ocr.tests.conftest import FakeEngine, make_png, make_tiff

pytestmark = pytest.mark.django_db

JOBS = "/api/ocr/jobs/"


def upload(client, files, language="auto", capture=None):
    payload = {"files[]": files, "language": language}
    if capture is None:
        return client.post(JOBS, payload, format="multipart")
    with capture(execute=True):
        return client.post(JOBS, payload, format="multipart")


def test_health(client):
    response = client.get("/api/health/")
    assert response.status_code == 200
    body = response.json()
    assert body["success"] is True and body["data"]["database"] == "ok"


def test_create_job_single_document(client, png_file, fake_engine, django_capture_on_commit_callbacks):
    response = upload(client, [png_file], "kk", django_capture_on_commit_callbacks)
    assert response.status_code == 201
    data = response.json()["data"]
    assert data["language"] == "kk" and data["documents_count"] == 1
    assert data["documents"][0]["status"] == "pending"  # snapshot taken at creation time

    job = OCRJob.objects.get(pk=data["id"])
    assert job.status == JobStatus.COMPLETED
    document = job.documents.get()
    assert document.status == DocumentStatus.COMPLETED
    assert document.pages_processed == 1 and document.processing_seconds is not None
    assert "ӘҒҚҢӨҰҮҺІ" in document.recognized_text


def test_two_documents_get_two_docx_and_zip(client, pdf_file, png_file, fake_engine, django_capture_on_commit_callbacks):
    response = upload(client, [pdf_file, png_file], "ru", django_capture_on_commit_callbacks)
    assert response.status_code == 201
    job_id = response.json()["data"]["id"]

    detail = client.get(f"{JOBS}{job_id}/").json()["data"]
    assert detail["status"] == "completed" and detail["zip_url"]
    assert [d["download_url"] is not None for d in detail["documents"]] == [True, True]

    # multi-page PDF -> one DOCX with both pages and a real table
    pdf_doc = next(d for d in detail["documents"] if d["original_filename"] == "report.pdf")
    assert pdf_doc["page_count"] == 2 and pdf_doc["pages_processed"] == 2
    download = client.get(pdf_doc["download_url"])
    assert download.status_code == 200
    assert "report.docx" in download["Content-Disposition"]
    docx = Document(io.BytesIO(b"".join(download.streaming_content)))
    texts = [p.text for p in docx.paragraphs]
    assert any("page 1" in t for t in texts) and any("page 2" in t for t in texts)
    assert len(docx.tables) == 2

    archive = client.get(detail["zip_url"])
    assert archive.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(b"".join(archive.streaming_content))).namelist()
    assert sorted(names) == ["report.docx", "scan.docx"]


def test_tiff_multipage_counts_pages(client, fake_engine, django_capture_on_commit_callbacks):
    tiff = SimpleUploadedFile("multi.tiff", make_tiff(3), content_type="image/tiff")
    response = upload(client, [tiff], "auto", django_capture_on_commit_callbacks)
    assert response.status_code == 201
    document = OCRJob.objects.get().documents.get()
    assert document.page_count == 3 and fake_engine.calls == 3


def test_rejects_unsupported_format(client):
    bad = SimpleUploadedFile("notes.txt", b"hello", content_type="text/plain")
    response = upload(client, [bad])
    assert response.status_code == 400
    body = response.json()
    assert body["success"] is False and "unsupported" in body["error"]["message"].lower()
    assert OCRJob.objects.count() == 0


def test_rejects_spoofed_content(client):
    fake = SimpleUploadedFile("evil.pdf", b"MZ not a pdf at all", content_type="application/pdf")
    response = upload(client, [fake])
    assert response.status_code == 400
    assert "does not match" in response.json()["error"]["message"]


def test_rejects_corrupted_image(client):
    broken = SimpleUploadedFile("broken.png", make_png()[:40], content_type="image/png")
    response = upload(client, [broken])
    assert response.status_code == 400


def test_rejects_too_many_files(client):
    files = [SimpleUploadedFile(f"{i}.png", make_png(), content_type="image/png") for i in range(3)]
    response = upload(client, files)
    assert response.status_code == 400
    assert "At most 2" in response.json()["error"]["message"]


def test_rejects_no_files_and_bad_language(client, png_file):
    assert client.post(JOBS, {}, format="multipart").status_code == 400
    response = upload(client, [png_file], "klingon")
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "validation_error"


def test_rejects_oversized_file(client, settings, png_file):
    settings.OCR_MAX_FILE_SIZE_MB = 0
    response = upload(client, [png_file])
    assert response.status_code == 400
    assert "larger than" in response.json()["error"]["message"]


def test_rejects_too_many_pages(client, settings):
    settings.OCR_MAX_PAGES_PER_DOCUMENT = 2
    tiff = SimpleUploadedFile("long.tif", make_tiff(3), content_type="image/tiff")
    response = upload(client, [tiff])
    assert response.status_code == 400 and "exceeds" in response.json()["error"]["message"]


def test_filename_is_sanitised(client, fake_engine, django_capture_on_commit_callbacks):
    sneaky = SimpleUploadedFile("../../etc/passwd.png", make_png(), content_type="image/png")
    response = upload(client, [sneaky], capture=django_capture_on_commit_callbacks)
    assert response.status_code == 201
    document = OCRJob.objects.get().documents.get()
    assert document.original_filename == "passwd.png"
    assert ".." not in document.source_file.name


def test_engine_unavailable_marks_job_failed(client, png_file, django_capture_on_commit_callbacks):
    # test settings use OCR_ENGINE_BACKEND=unavailable, so the real engine factory reports a clear error
    response = upload(client, [png_file], capture=django_capture_on_commit_callbacks)
    assert response.status_code == 201
    job = OCRJob.objects.get()
    assert job.status == JobStatus.FAILED
    document = job.documents.get()
    assert document.status == DocumentStatus.FAILED and "disabled" in document.error_message
    assert client.get(f"{JOBS}{job.id}/documents/{document.id}/download/").status_code == 409
    assert client.get(f"{JOBS}{job.id}/download-zip/").status_code == 409


def test_partial_failure(client, monkeypatch, django_capture_on_commit_callbacks):
    from apps.ocr.services import ocr_engine

    engine = FakeEngine()
    original = engine.recognize_page
    state = {"n": 0}

    def flaky(path, language):
        state["n"] += 1
        if state["n"] == 2:
            raise RuntimeError("boom")
        return original(path, language)

    engine.recognize_page = flaky
    monkeypatch.setattr(ocr_engine, "get_engine", lambda: engine)
    files = [SimpleUploadedFile(f"{i}.png", make_png(), content_type="image/png") for i in range(2)]
    upload(client, files, capture=django_capture_on_commit_callbacks)
    job = OCRJob.objects.get()
    assert job.status == JobStatus.PARTIALLY_COMPLETED
    statuses = sorted(job.documents.values_list("status", flat=True))
    assert statuses == ["completed", "failed"]
    failed = job.documents.get(status="failed")
    assert "RuntimeError" in failed.error_message and "boom" not in failed.error_message


def test_list_retrieve_delete(client, png_file, fake_engine, django_capture_on_commit_callbacks):
    job_id = upload(client, [png_file], capture=django_capture_on_commit_callbacks).json()["data"]["id"]
    listing = client.get(JOBS).json()["data"]
    assert [j["id"] for j in listing] == [job_id]

    document = OCRJob.objects.get().documents.get()
    path = document.source_file.path
    assert client.delete(f"{JOBS}{job_id}/").status_code == 200
    assert client.get(f"{JOBS}{job_id}/").status_code == 404
    import os

    assert not os.path.exists(path)


def test_unknown_job_is_404_in_envelope(client):
    response = client.get(f"{JOBS}00000000-0000-0000-0000-000000000000/")
    assert response.status_code == 404
    assert response.json()["success"] is False


def test_queue_failure_marks_job_failed(client, png_file, monkeypatch, django_capture_on_commit_callbacks):
    from apps.ocr import views

    def broken_delay(job_id):
        raise ConnectionError("redis down")

    monkeypatch.setattr(views.process_ocr_job, "delay", broken_delay)
    upload(client, [png_file], capture=django_capture_on_commit_callbacks)
    job = OCRJob.objects.get()
    assert job.status == JobStatus.FAILED and "queue" in job.error_message
